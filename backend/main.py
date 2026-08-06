import json

from fastapi import FastAPI,HTTPException,Query
from fastapi.responses import FileResponse
import mysql.connector as ms
from fastapi.middleware.cors import CORSMiddleware
from crud import insert_into_records
from models import StockItem,ReturnedItem
from services import add_to_stores, clear_temp_data, get_code_details, is_valid_name, lookupRange, make_valid_table_name, ordinal, reverse_table_name, submit_new_stock,add_design_temp,temp_stock_data,from_shelf,lookup,add_design_temp_return,submit_returned_stock,remove_from_temp,lookupforprint,submit_sales_stock, safe_identifier
from database import get_db_connection
from datetime import datetime
from openpyxl.styles import Font, Alignment
from fastapi.responses import StreamingResponse, JSONResponse
import io
import pandas as pd
from openpyxl.styles import Border, Side
from typing import Optional

app = FastAPI()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:5173"],
    allow_credentials=True,
    allow_methods=["*"],  
    allow_headers=["*"],  
    expose_headers=["Content-Disposition"],
)

#health check
@app.get("/health/db")
async def check_db_health():
	try:
		connection = get_db_connection()
		connection.ping(reconnect=True)
		connection.close()
		return {"mysql": "ok"}
	except Exception as e:
		return JSONResponse(status_code=500, content={"mysql": "down", "error": str(e)})
	

#rename brandname
@app.post("/alter/brandname")
async def alterName(
	old_name : str = Query(...),
	new_name : str = Query(...)
):
	old_name = make_valid_table_name(old_name)
	new_name = make_valid_table_name(new_name)

	if not is_valid_name(old_name) or not is_valid_name(new_name):
		raise HTTPException(status_code=400, detail="Brand name should have only letters, numbers and spaces.")

	try:
		connection = get_db_connection()
		cursor = connection.cursor()

		cursor.execute(f"CREATE DATABASE IF NOT EXISTS `{new_name}`;")
		cursor.execute(f"""
					SELECT table_name 
					FROM information_schema.tables 
					WHERE table_schema = '{old_name}';
			""")
		tables = cursor.fetchall()
		for each in tables:
			print(each)

		for (table_name,) in tables:
				cursor.execute(f"""
						RENAME TABLE `{old_name}`.`{table_name}` TO `{new_name}`.`{table_name}`;
				""")

		cursor.execute(f"DROP DATABASE `{old_name}`;")

		connection.commit()
		cursor.close()

		connection.close()

		return {
			"message" : f"BrandName changed from `{reverse_table_name(old_name)}` to `{reverse_table_name(new_name)}` successfully",
			"new_name" : reverse_table_name(new_name)
		}		
	except Exception as e:
		raise HTTPException(status_code=400, detail=str(e))	

#get brandnames and store names
@app.get("/brands")
async def get_brands():
	try:
		connection = get_db_connection()
		cursor = connection.cursor()
		
		cursor.execute("""
            SELECT schema_name 
            FROM information_schema.schemata 
            WHERE schema_name NOT IN ('mysql', 'information_schema','performance_schema', 'sys');
        """)
		brands = [row[0] for row in cursor.fetchall()]
		for i in range(0,len(brands)):
			brands[i] = reverse_table_name(brands[i])

		connection.close()
		return {"brands": brands}
	except Exception as e:
		raise HTTPException(status_code=400, detail=str(e))

@app.get("/stores")
async def get_stores(
	brand_name : str
):
	try:
		connection = get_db_connection(brand_name)
		cursor = connection.cursor()
		
		cursor.execute("SELECT store_name FROM stores;")

		thestores = [row[0] for row in cursor.fetchall()]
		for i in range(0,len(thestores)):
			thestores[i] = reverse_table_name(thestores[i])

		connection.close()
		return {"stores": thestores}
	except Exception as e:
		raise HTTPException(status_code=400, detail=str(e))		

#add new items
@app.post("/add/new")
async def add_stock_item(
	stock_item : StockItem,
	store_name : str = Query(...), 
	date : str = Query(...)
	):
	
	date_obj = datetime.strptime(date, "%Y-%m-%d")
	formatted_date = date_obj.strftime("%d_%b_%Y")

	store_name = safe_identifier(make_valid_table_name(store_name))
	store_key = f"{store_name}_{formatted_date}_new_stock"

	try:
		temp_data = add_design_temp(store_key,stock_item)
		return {
			"store_key" : store_key,
			"current_designs" : temp_data
		}

	except Exception as e:
		raise HTTPException(status_code=400, detail=str(e))

#returned
@app.post("/add/return")
async def add_returned_item(
	returned_item : ReturnedItem,
	store_name : str = Query(...), 
	date : str = Query(...)
	):
	
	date_obj = datetime.strptime(date, "%Y-%m-%d")
	formatted_date = date_obj.strftime("%d_%b_%Y")
	
	store_name = safe_identifier(make_valid_table_name(store_name))
	store_key = f"{store_name}_{formatted_date}_return_stock"

	try:
		temp_data = add_design_temp_return(store_key,returned_item)
		return {
			"message" : "Stock item added to returned list",
			"store_key" : store_key,
			"current_designs" : temp_data
		}

	except Exception as e:
		raise HTTPException(status_code=400, detail=str(e))

#sales
@app.post("/add/sales")
async def add_sales_item(
	returned_item : ReturnedItem,
	store_name : str = Query(...), 
	date : str = Query(...)
	):
	
	date_obj = datetime.strptime(date, "%Y-%m-%d")
	formatted_date = date_obj.strftime("%d_%b_%Y")
	
	store_name = safe_identifier(make_valid_table_name(store_name))
	store_key = f"{store_name}_{formatted_date}_sales_stock"

	try:
		temp_data = add_design_temp_return(store_key,returned_item)
		return {
			"message" : "Stock item added to returned list",
			"store_key" : store_key,
			"current_designs" : temp_data
		}

	except Exception as e:
		raise HTTPException(status_code=400, detail=str(e))

@app.get("/view/{store_key}")
async def view_stock(store_key: str):
    if store_key not in temp_stock_data or not temp_stock_data[store_key]:
        raise HTTPException(status_code=404, detail="Add designs.")
    
    return {"message": "Temporary stock items", "data": temp_stock_data[store_key]}	

#view shelf
@app.get("/shelf/{brand_name}/{store_name}")
async def view_shelf(brand_name : str, store_name : str):
	store_name = safe_identifier(make_valid_table_name(store_name))
	brand_name = safe_identifier(make_valid_table_name(brand_name))

	try:
		connection = get_db_connection(brand_name)
		shelf = from_shelf(store_name, connection)
		connection.close()

		if not shelf:
			return {"message": f"No stock found for {store_name}.", "data": []}

		return {"message": f"Stock at {store_name}.", "data": shelf}
	
	except Exception as e:
		raise HTTPException(status_code=500, detail=str(e))

#view code.
@app.get("/view_code/{brand_name}/{design_code}")
async def view_code(brand_name: str, design_code: str):
    brand_name = make_valid_table_name(brand_name)
    design_code = design_code.strip()

    try:
        connection = get_db_connection(brand_name)
        code_data = get_code_details(design_code, connection)
        connection.close()

        if not code_data:
            return {
                "message": f"Design code '{design_code}' not found.",
                "data": {}
            }

        return {
            "message": f"Stock locations for design code '{design_code}'.",
            "data": code_data
        }

    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

#view action on a date
@app.get("/view_action/{brand_name}/{store_name}/{date}/{action}")
async def view_action(brand_name : str,store_name : str, date: str, action : str):
	store_name = safe_identifier(make_valid_table_name(store_name))
	brand_name = safe_identifier(make_valid_table_name(brand_name))

	try:
		connection = get_db_connection(brand_name)
		action_on_date = lookup(store_name,date,action,connection)
		connection.close()

		if not action_on_date:
			return {"message": f"Action not performed at {store_name} on {date}. ", "data" : []}

		return {"message": f"{action} stock on {date} at {store_name}.", "data": action_on_date}
	
	except Exception as e:
		raise HTTPException(status_code=500, detail=str(e))
	
#view records in date range
@app.get("/view_range/{brand_name}")
async def view_range(
    brand_name: str,
    fromDate: str,
    toDate: str,
    store_name: Optional[str] = Query(None),
    action: Optional[str] = Query(None)
):
    brand_name = safe_identifier(make_valid_table_name(brand_name))

    try:
        connection = get_db_connection(brand_name)
        action_on_range= lookupRange(fromDate, toDate, store_name, action, connection)
        connection.close()

        if not action_on_range:
            return {
                "message": f"No actions performed from {fromDate} to {toDate}.",
                "data": []
            }

        return {
            "message": f"Records from {fromDate} to {toDate}",
            "data": action_on_range
        }

    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@app.post("/submit/{brand_name}/{action}")
async def submition_handler(
	brand_name : str,
	action : str,
	store_name : str = Query(...), 
	date : str = Query(...)
	):

	brand_name = safe_identifier(make_valid_table_name(brand_name))
	store_name = safe_identifier(make_valid_table_name(store_name))

	date_obj = datetime.strptime(date, "%Y-%m-%d")
	formatted_date = date_obj.strftime("%d_%b_%Y")
	
	table_name = f"{store_name}_{formatted_date}_{action}"
	connection = None
	cursor = None

	try:
		connection = get_db_connection(brand_name)
        
		add_to_stores(store_name,connection)
		
		cursor = connection.cursor()
            
		insert_into_records(cursor,reverse_table_name(store_name),action,date)
		cursor.close()
		cursor = None

		if(action == "new"):
			store_key = f"{store_name}_{formatted_date}_new_stock"
			table_name = f"{store_name}_{formatted_date}_new_stock"
			submit_new_stock(store_name, store_key,table_name,connection)
		elif action == "return":
			store_key = f"{store_name}_{formatted_date}_return_stock"
			table_name = f"{store_name}_{formatted_date}_return_stock"
			submit_returned_stock(store_name, store_key,table_name,formatted_date,connection)
		elif action == "sales":
			store_key = f"{store_name}_{formatted_date}_sales_stock"
			table_name = f"{store_name}_{formatted_date}_sales_stock"
			submit_sales_stock(store_name, store_key,table_name,formatted_date,connection)	

		connection.commit()
            
		del temp_stock_data[store_key]
		
		connection.close()

		return {
			"message" : "Stock details submitted successfully"
		}

	except Exception as e:
		if connection:
			connection.rollback()
		raise HTTPException(status_code=500, detail=str(e))

	finally:
		if cursor:
			cursor.close()
		if connection:
			connection.close()

#remove from temp
@app.post("/remove/temp/{code}")
async def remove(
	code : str,
	store_key : str = Query(...)
	):
	
	temp_data = remove_from_temp(store_key,code)
	
	return {
		"message" : "Updated added list",
		"data" : temp_data
	}

@app.post("/clear/temp")
async def clearTemp(store_key: str = Query(...)):
    temp_data = clear_temp_data(store_key)
    return {
        "message": f"Cleared temp data for store '{store_key}'.",
        "data": temp_data,
    }

#select columns
@app.get("/columns/{brand_name}")
async def get_columns(
    brand_name: str,
    store_name: str = Query(...),
    date: str = Query(...),
    action: str = Query(...)
):
    date_obj = datetime.strptime(date, "%Y-%m-%d")
    formatted_date = date_obj.strftime("%d_%b_%Y")

    brand_name = safe_identifier(make_valid_table_name(brand_name))
    store_name = safe_identifier(make_valid_table_name(store_name))

    if action == "new":
        table_name = f"{store_name}_{formatted_date}_new_stock"
    elif action == "return":
        table_name = f"{store_name}_{formatted_date}_return_stock"
    elif action == "sales":
        table_name = f"{store_name}_{formatted_date}_sales_stock"
    else:
        raise HTTPException(status_code=400, detail="Invalid action")

    try:
        connection = get_db_connection(brand_name)
        cursor = connection.cursor(dictionary=True)

        cursor.execute(f"SELECT * FROM `{table_name}` LIMIT 1")
        row = cursor.fetchone()
        cursor.close()
        connection.close()

        if not row:
            raise HTTPException(status_code=404, detail="No data found.")

        fixed_columns = []
        custom_field_keys = []
        per_size_custom_fields_keys = []
		
        for key, value in row.items():
            lower = key.lower()
            if lower == "custom_fields":
                # parse the dict and extract keys
                try:
                    cf = json.loads(value) if isinstance(value, str) else value
                    if isinstance(cf, dict):
                        custom_field_keys = list(cf.keys())
                except:
                    pass
            elif lower == "per_size_custom_fields":
                try:
                    pcf = json.loads(value) if isinstance(value, str) else value
                    if isinstance(pcf, dict):
                        per_size_custom_fields_keys = list(pcf.keys())
                except:
                    pass
            else:
                fixed_columns.append(key)

        return {
            "fixed_columns": fixed_columns,
            "custom_field_keys": custom_field_keys,
            "per_size_custom_fields_keys" : per_size_custom_fields_keys
        }

    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

#print excel
@app.post("/printExcel/{brand_name}")
async def print_table_excel(
    brand_name: str,
    store_name: str = Query(...),
    date: str = Query(...),
    action: str = Query(...),
    selected_columns: str = Query(default=None)  # comma-separated
):
    # Parse and format date
    date_obj = datetime.strptime(date, "%Y-%m-%d")
    formatted_date = date_obj.strftime("%d_%b_%Y")
    readable = f"{ordinal(date_obj.day)} {date_obj.strftime('%b')}, {date_obj.year}"

    # Sanitize brand and store names
    brand_name = safe_identifier(make_valid_table_name(brand_name))
    store_name = safe_identifier(make_valid_table_name(store_name))

    # Get data from database
    connection = get_db_connection(brand_name)
    stock_data = lookupforprint(store_name, date, action, connection)

    if not stock_data:
        raise HTTPException(status_code=404, detail="No stock data found.")

    # Flatten custom_fields into top-level keys
    for item in stock_data:
        if "custom_fields" in item and item["custom_fields"]:
            try:
                cf = (
                    json.loads(item["custom_fields"])
                    if isinstance(item["custom_fields"], str)
                    else item["custom_fields"]
                )
                if isinstance(cf, dict):
                    for k, v in cf.items():
                        item[k] = v
            except:
                pass
        if "per_size_custom_fields" in item and item["per_size_custom_fields"]:
            try:
                pcf = (
                    json.loads(item["per_size_custom_fields"])
                    if isinstance(item["per_size_custom_fields"], str)
                    else item["per_size_custom_fields"]
                )
                if isinstance(pcf, dict):
                    for k, v in pcf.items():
                        item[k] = v
            except:
                pass                

    # Define default columns
    all_possible = [
        ("item", "Item"),
        ("design_code", "Code"),
        ("hsn_code", "HSN Code"),
        ("size", "Size"),
        ("sp_per_item", "Price"),
        ("qty", "Qty"),
        ("gst_rate", "GST Rate"),
        ("taxable_amount", "Taxable"),
        ("tax_amount", "Tax"),
    ]

    # Add custom fields from first record
    first = stock_data[0]
    for key in first.keys():
        if key not in [c[0] for c in all_possible] and key != "custom_fields" and key != "per_size_custom_fields":
            all_possible.append((key, key.title()))

    # Select columns based on user input or GST presence
    if selected_columns:
        chosen = selected_columns.split(",")
        columns = [(k, label) for k, label in all_possible if k in chosen]
    else:
        has_gst = any(float(item.get("gst_rate", 0) or 0) > 0 for item in stock_data)
        columns = [
            c for c in all_possible if c[0] not in ("gst_rate", "taxable_amount", "tax_amount")
        ] if not has_gst else all_possible

    # Prepare DataFrame
    df = pd.DataFrame(stock_data)
    available = [c for c in columns if c[0] in df.columns]
    df = df[[c[0] for c in available]]
    df.columns = [c[1] for c in available]
    df.index = range(1, len(df) + 1)

    # Write Excel to BytesIO
    excel_bytes_io = io.BytesIO()
    with pd.ExcelWriter(excel_bytes_io, engine="openpyxl") as writer:
        sheet_name = "Stock Data"
        df.to_excel(writer, index=True, startrow=3, sheet_name=sheet_name, index_label="S.No")

        workbook = writer.book
        worksheet = writer.sheets[sheet_name]

        # Add headers
        worksheet["C1"] = brand_name.title()
        worksheet["D1"] = store_name.title()
        worksheet["E2"] = readable
        worksheet["F2"] = f"{action.title()} Stocks"

        for cell_ref in ["C1", "D1"]:
            worksheet[cell_ref].alignment = Alignment(horizontal="center", vertical="center")

        # Set column widths
        for i, col_letter in enumerate("ABCDEFGHIJKLMNOP"[:len(available) + 1]):
            worksheet.column_dimensions[col_letter].width = 16

        # Apply border and alignment
        thin_border = Border(
            left=Side(style="thin", color="000000"),
            right=Side(style="thin", color="000000"),
            top=Side(style="thin", color="000000"),
            bottom=Side(style="thin", color="000000")
        )

        for row in worksheet.iter_rows(
            min_row=1, max_row=worksheet.max_row, min_col=1, max_col=len(available) + 1
        ):
            for cell in row:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = thin_border

    excel_bytes_io.seek(0)

    return StreamingResponse(
        excel_bytes_io,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="{reverse_table_name(store_name).title()} {readable} {action.title()} stocks.xlsx"'
        }
    )

if __name__ == "__main__":
    import uvicorn
    print("Running FastAPI server...")
    uvicorn.run(app, host="127.0.0.1", port=8000)